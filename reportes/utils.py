from django.http import HttpResponse, HttpResponseBadRequest
from django.template.loader import render_to_string, TemplateDoesNotExist
from datetime import datetime
import io 
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill 
from openpyxl.utils import get_column_letter
import logging 

logger = logging.getLogger(__name__)

from accounts.models import Usuario
from gestion_usuarios.models import Rol
from gestion_clientes.models import Cliente 
from gestion_servicios.models import Servicio 
from gestion_especialistas.models import Especialista 
from gestion_citas.models import Cita 

try:
    from weasyprint import HTML
except ImportError:
    HTML = None 


def _filtrar_usuarios(fecha_inicial, fecha_final, rol_id):
    usuarios = Usuario.objects.all().order_by('id')
    
    if fecha_inicial:
        fecha_ini_dt = datetime.strptime(fecha_inicial, '%Y-%m-%d').date()
        usuarios = usuarios.filter(date_joined__date__gte=fecha_ini_dt) 
    
    if fecha_final:
        fecha_fin_dt = datetime.strptime(fecha_final, '%Y-%m-%d').date()
        usuarios = usuarios.filter(date_joined__date__lte=fecha_fin_dt)
        
    if rol_id:
        usuarios = usuarios.filter(rol__id=rol_id) 
        
    return usuarios

def _filtrar_clientes(fecha_inicial, fecha_final):
    clientes = Cliente.objects.all().order_by('id')
    
    if fecha_inicial:
        fecha_ini_dt = datetime.strptime(fecha_inicial, '%Y-%m-%d').date()
        clientes = clientes.filter(fecha_creacion__date__gte=fecha_ini_dt)
    
    if fecha_final:
        fecha_fin_dt = datetime.strptime(fecha_final, '%Y-%m-%d').date()
        clientes = clientes.filter(fecha_creacion__date__lte=fecha_fin_dt)
        
    return clientes

def _filtrar_servicios():
    servicios = Servicio.objects.all().order_by('nombre')
    return servicios

def _filtrar_especialistas(fecha_inicial, fecha_final):
    especialistas = Especialista.objects.all().select_related('usuario').prefetch_related('servicios').order_by('id')
    return especialistas

# =================================================================
# CORRECCIÓN 1/2: Optimización de filtrado de Citas
# Eliminamos la precarga anidada (especialista__usuario) ya que 
# cita.especialista ya es un objeto Usuario.
# =================================================================
def _filtrar_citas(fecha_inicial, fecha_final):
    # La precarga es correcta: cliente, servicio, especialista (especialista ya es el Usuario)
    citas = Cita.objects.all().select_related('cliente', 'servicio', 'especialista').order_by('fecha_hora')
    
    if fecha_inicial:
        fecha_ini_dt = datetime.strptime(fecha_inicial, '%Y-%m-%d').date()
        citas = citas.filter(fecha_hora__date__gte=fecha_ini_dt)
    
    if fecha_final:
        fecha_fin_dt = datetime.strptime(fecha_final, '%Y-%m-%d').date()
        citas = citas.filter(fecha_hora__date__lte=fecha_fin_dt)
        
    return citas


def _get_estilos_excel():
    titulo_font = Font(bold=True, size=14)
    encabezado_font = Font(bold=True) 
    center_align = Alignment(horizontal='center', vertical='center')
    borde_fino = Side(border_style="thin", color="000000")
    borde_completo = Border(top=borde_fino, bottom=borde_fino, left=borde_fino, right=borde_fino) 
    relleno_gris_claro = PatternFill(start_color="BFBFBF", end_color="BFBFBF", fill_type="solid") 
    return titulo_font, encabezado_font, center_align, borde_completo, relleno_gris_claro

# =================================================================
# CORRECCIÓN 2/2: Lógica de procesamiento de Citas
# Accedemos directamente a first_name/last_name desde cita.especialista.
# =================================================================
def _procesar_citas(citas_qs):
    citas_procesadas = []
    for cita in citas_qs:
        
        nombre_especialista = "No Asignado"
        # CORRECCIÓN APLICADA AQUÍ: cita.especialista YA ES el objeto Usuario
        if cita.especialista:
            usuario = cita.especialista # Renombramos para claridad, es el objeto Usuario
            nombre_especialista = f"{usuario.first_name} {usuario.last_name}".strip()
            # Si los campos están vacíos, usamos el username como fallback
            if not nombre_especialista:
                 nombre_especialista = usuario.username
        
        nombre_cliente = "Cliente Eliminado"
        if hasattr(cita, 'cliente') and cita.cliente:
            try:
                # Intentamos acceder a los campos directos del modelo Cliente
                nombre_cliente = f"{cita.cliente.nombre} {cita.cliente.apellido}"
            except AttributeError:
                # Si falla, es probable que cita.cliente sea un objeto Usuario
                # o que el Cliente redirija a Usuario.
                # Intentamos obtener el nombre del Usuario si existe
                if hasattr(cita.cliente, 'usuario') and cita.cliente.usuario:
                    nombre_cliente = f"{cita.cliente.usuario.first_name} {cita.cliente.usuario.last_name}".strip()
                else:
                    # Si no podemos resolverlo, usamos el full name del cliente (si es un Usuario)
                    # o un placeholder seguro
                    if hasattr(cita.cliente, 'get_full_name'):
                        nombre_cliente = cita.cliente.get_full_name().strip()
                    else:
                        nombre_cliente = "Cliente Inaccesible"
                
                # Fallback: Si el cliente es un Usuario y sigue vacío
                if not nombre_cliente and hasattr(cita.cliente, 'username'):
                    nombre_cliente = cita.cliente.username

        citas_procesadas.append({
            'cita': cita,
            'nombre_especialista_resuelto': nombre_especialista,
            'nombre_cliente_resuelto': nombre_cliente
        })
    return citas_procesadas


def generar_excel_usuarios(fecha_inicial=None, fecha_final=None, rol_id=None):
    usuarios = _filtrar_usuarios(fecha_inicial, fecha_final, rol_id)

    wb = Workbook()
    ws = wb.active
    
    titulo_font, encabezado_font, center_align, borde_completo, relleno_gris_claro = _get_estilos_excel()
    
    ws.merge_cells('A1:E1') 
    ws['A1'] = 'Reporte de Usuarios y Roles'
    ws['A1'].font = titulo_font
    ws['A1'].alignment = center_align

    ws.merge_cells('A2:E2')
    ws['A2'] = f"Fecha de Generación: {datetime.now().strftime('%d/%m/%Y %H:%M')}" 
    ws['A2'].alignment = center_align

    columnas = ['ID', 'Nombre Completo', 'Email', 'Roles Asignados', 'Fecha Creacion']
    ancho_columnas = {'A': 5, 'B': 30, 'C': 35, 'D': 20, 'E': 25}

    for col_num, column_title in enumerate(columnas, 1):
        cell_letter = get_column_letter(col_num)
        cell = ws.cell(row=4, column=col_num, value=column_title) 
        cell.font = encabezado_font
        cell.fill = relleno_gris_claro
        cell.border = borde_completo
        ws.column_dimensions[cell_letter].width = ancho_columnas.get(cell_letter, 15)
    
    fila_actual = 5
    for usuario in usuarios:
        nombre_completo = usuario.get_full_name()
        if not nombre_completo:
            nombre_completo = usuario.username

        roles_asignados = "N/A"
        if hasattr(usuario, 'rol') and usuario.rol: 
            roles_asignados = usuario.rol.nombre
        
        datos_fila = [
            usuario.id, 
            nombre_completo,
            usuario.email, 
            roles_asignados,
            usuario.date_joined.strftime('%Y-%m-%d %H:%M:%S') if usuario.date_joined else '', 
        ]
        
        ws.append(datos_fila)

        for col_num in range(1, len(columnas) + 1):
            cell = ws.cell(row=fila_actual, column=col_num)
            cell.border = borde_completo

        fila_actual += 1

    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    filename = f"reporte_usuarios_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    response['Content-Disposition'] = f'attachment; filename="{filename}"'

    wb.save(response)

    return response

def generar_excel_clientes(fecha_inicial=None, fecha_final=None):
    clientes = _filtrar_clientes(fecha_inicial, fecha_final)

    wb = Workbook()
    ws = wb.active
    
    titulo_font, encabezado_font, center_align, borde_completo, relleno_gris_claro = _get_estilos_excel()
    
    ws.merge_cells('A1:G1') 
    ws['A1'] = 'Reporte de Clientes Registrados'
    ws['A1'].font = titulo_font
    ws['A1'].alignment = center_align

    ws.merge_cells('A2:G2')
    ws['A2'] = f"Fecha de Generación: {datetime.now().strftime('%d/%m/%Y %H:%M')}" 
    ws['A2'].alignment = center_align
    
    columnas = ['ID', 'Nombre', 'Apellido', 'Teléfono', 'Email', 'Fecha Registro', 'Activo']
    ancho_columnas = {'A': 5, 'B': 20, 'C': 20, 'D': 15, 'E': 35, 'F': 20, 'G': 10}

    for col_num, column_title in enumerate(columnas, 1):
        cell_letter = get_column_letter(col_num)
        cell = ws.cell(row=4, column=col_num, value=column_title) 
        
        cell.font = encabezado_font
        cell.fill = relleno_gris_claro
        cell.border = borde_completo
        
        ws.column_dimensions[cell_letter].width = ancho_columnas.get(cell_letter, 15)
    
    fila_actual = 5
    for cliente in clientes:
        datos_fila = [
            cliente.id, 
            cliente.nombre,
            cliente.apellido,
            cliente.telefono, 
            cliente.email, 
            cliente.fecha_creacion.strftime('%Y-%m-%d %H:%M:%S') if cliente.fecha_creacion else '',
            'Sí' if cliente.is_activo else 'No'
        ]
        
        ws.append(datos_fila)

        for col_num in range(1, len(columnas) + 1):
            cell = ws.cell(row=fila_actual, column=col_num)
            cell.border = borde_completo

        fila_actual += 1

    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    filename = f"reporte_clientes_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    response['Content-Disposition'] = f'attachment; filename="{filename}"'

    wb.save(response)

    return response


def generar_excel_servicios():
    servicios = _filtrar_servicios()

    wb = Workbook()
    ws = wb.active
    
    titulo_font, encabezado_font, center_align, borde_completo, relleno_gris_claro = _get_estilos_excel()
    
    ws.merge_cells('A1:E1') 
    ws['A1'] = 'Reporte de Servicios Disponibles'
    ws['A1'].font = titulo_font
    ws['A1'].alignment = center_align

    ws.merge_cells('A2:E2')
    ws['A2'] = f"Fecha de Generación: {datetime.now().strftime('%d/%m/%Y %H:%M')}" 
    ws['A2'].alignment = center_align

    columnas = ['ID', 'Nombre del Servicio', 'Precio Base ($)', 'Duración (min)', 'Estado']
    ancho_columnas = {'A': 5, 'B': 30, 'C': 15, 'D': 15, 'E': 10}

    for col_num, column_title in enumerate(columnas, 1):
        cell_letter = get_column_letter(col_num)
        cell = ws.cell(row=4, column=col_num, value=column_title) 
        cell.font = encabezado_font
        cell.fill = relleno_gris_claro
        cell.border = borde_completo
        ws.column_dimensions[cell_letter].width = ancho_columnas.get(cell_letter, 15)
    
    fila_actual = 5
    for servicio in servicios:
        
        ws.cell(row=fila_actual, column=1, value=servicio.id).border = borde_completo
        ws.cell(row=fila_actual, column=2, value=servicio.nombre).border = borde_completo
        
        precio_cell = ws.cell(row=fila_actual, column=3, value=servicio.precio) 
        precio_cell.number_format = '"$ "#,##0.00' 
        precio_cell.border = borde_completo
        
        ws.cell(row=fila_actual, column=4, value=servicio.duracion_minutos).border = borde_completo
        
        ws.cell(row=fila_actual, column=5, value=servicio.estado).border = borde_completo 

        fila_actual += 1

    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    filename = f"reporte_servicios_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    response['Content-Disposition'] = f'attachment; filename="{filename}"'

    wb.save(response)

    return response


def generar_excel_especialistas(fecha_inicial=None, fecha_final=None):
    especialistas = _filtrar_especialistas(fecha_inicial, fecha_final)

    wb = Workbook()
    ws = wb.active
    
    titulo_font, encabezado_font, center_align, borde_completo, relleno_gris_claro = _get_estilos_excel()
    
    ws.merge_cells('A1:F1') 
    ws['A1'] = 'Reporte de Especialistas (Empleados)'
    ws['A1'].font = titulo_font
    ws['A1'].alignment = center_align

    ws.merge_cells('A2:F2')
    ws['A2'] = f"Fecha de Generación: {datetime.now().strftime('%d/%m/%Y %H:%M')}" 
    ws['A2'].alignment = center_align

    columnas = ['ID', 'Nombre Completo', 'Email', 'Especialidad', 'Servicios que realiza', 'Activo'] 
    ancho_columnas = {'A': 5, 'B': 30, 'C': 35, 'D': 25, 'E': 45, 'F': 10} 
    
    for col_num, column_title in enumerate(columnas, 1):
        cell_letter = get_column_letter(col_num)
        cell = ws.cell(row=4, column=col_num, value=column_title) 
        
        cell.font = encabezado_font
        cell.fill = relleno_gris_claro
        cell.border = borde_completo
        
        ws.column_dimensions[cell_letter].width = ancho_columnas.get(cell_letter, 15)
    
    fila_actual = 5
    for esp in especialistas:
        usuario = esp.usuario if hasattr(esp, 'usuario') and esp.usuario else None
        
        nombre_completo = f"{usuario.first_name} {usuario.last_name}" if usuario else 'N/A'
        email = usuario.email if usuario else 'N/A'
        
        servicios_realiza = ', '.join([s.nombre for s in esp.servicios.all()])
        
        datos_fila = [
            esp.id, 
            nombre_completo,
            email,
            esp.especialidad, 
            servicios_realiza, 
            'Sí' if esp.is_activo else 'No'
        ]
        
        ws.append(datos_fila)

        for col_num in range(1, len(columnas) + 1):
            cell = ws.cell(row=fila_actual, column=col_num)
            cell.border = borde_completo

        fila_actual += 1

    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    filename = f"reporte_especialistas_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    response['Content-Disposition'] = f'attachment; filename="{filename}"'

    wb.save(response)

    return response

def generar_excel_citas(fecha_inicial=None, fecha_final=None):
    citas_qs = _filtrar_citas(fecha_inicial, fecha_final)
    citas_procesadas = _procesar_citas(citas_qs)

    wb = Workbook()
    ws = wb.active
    
    titulo_font, encabezado_font, center_align, borde_completo, relleno_gris_claro = _get_estilos_excel()
    
    ws.merge_cells('A1:F1') 
    ws['A1'] = 'Reporte de Citas Agendadas'
    ws['A1'].font = titulo_font
    ws['A1'].alignment = center_align

    ws.merge_cells('A2:F2')
    ws['A2'] = f"Fecha de Generación: {datetime.now().strftime('%d/%m/%Y %H:%M')}" 
    ws['A2'].alignment = center_align

    columnas = ['ID', 'Cliente', 'Servicio', 'Especialista', 'Fecha y Hora', 'Estado']
    ancho_columnas = {'A': 5, 'B': 30, 'C': 25, 'D': 30, 'E': 25, 'F': 15}

    for col_num, column_title in enumerate(columnas, 1):
        cell_letter = get_column_letter(col_num)
        cell = ws.cell(row=4, column=col_num, value=column_title) 
        cell.font = encabezado_font
        cell.fill = relleno_gris_claro
        cell.border = borde_completo
        ws.column_dimensions[cell_letter].width = ancho_columnas.get(cell_letter, 15)
    
    fila_actual = 5
    
    for item in citas_procesadas:
        cita = item['cita']
        
        datos_fila = [
            cita.id, 
            item['nombre_cliente_resuelto'],
            cita.servicio.nombre if cita.servicio else 'N/A',
            item['nombre_especialista_resuelto'],
            cita.fecha_hora.strftime('%Y-%m-%d %H:%M'), 
            cita.get_estado_display(),
        ]
        
        ws.append(datos_fila)

        for col_num in range(1, len(columnas) + 1):
            cell = ws.cell(row=fila_actual, column=col_num)
            cell.border = borde_completo

        fila_actual += 1

    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    filename = f"reporte_citas_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    response['Content-Disposition'] = f'attachment; filename="{filename}"'

    wb.save(response)

    return response


def generar_pdf_usuarios(fecha_inicial=None, fecha_final=None, rol_id=None, formato='pdf'):
    if HTML is None:
        return HttpResponse("Error 501: La conversión a PDF binario (WeasyPrint) no está disponible.", status=501)
    
    usuarios = _filtrar_usuarios(fecha_inicial, fecha_final, rol_id)
        
    rol_filtrado_nombre = "Todos los Roles"
    if rol_id:
        try:
            rol_filtrado_nombre = Rol.objects.get(id=rol_id).nombre
        except Rol.DoesNotExist:
            pass 
        
    context = {
        'titulo': 'Reporte de Usuarios y Roles',
        'fecha_generacion': datetime.now().strftime('%d/%m/%Y %H:%M:%S'),
        'usuarios': usuarios,
        'filtro_fecha_inicial': fecha_inicial,
        'filtro_fecha_final': fecha_final,
        'filtro_rol_nombre': rol_filtrado_nombre,
    }
    
    try:
        html_string = render_to_string('reportes/reporte_usuarios_pdf.html', context)
        html_doc = HTML(string=html_string)
        pdf_file = html_doc.write_pdf()

        response = HttpResponse(pdf_file, content_type='application/pdf')
        filename = f"reporte_usuarios_{datetime.now().strftime('%Y%m%d_%H%M%S')}.pdf"
        
        disposition = 'inline' if formato in ['pdf_imprimir', 'inline'] else 'attachment'
        response['Content-Disposition'] = f'{disposition}; filename="{filename}"' 
            
        return response
    except TemplateDoesNotExist:
        logger.error("Error fatal: No se encontró la plantilla 'reportes/reporte_usuarios_pdf.html'")
        return HttpResponse("Error 500: No se encontró la plantilla 'reportes/reporte_usuarios_pdf.html'. Verifique la ruta.", status=500)
    except Exception as e:
        logger.error(f"Error al generar PDF de Usuarios: {e}")
        return HttpResponse(f"Error al generar PDF: {e}", status=500)


def generar_pdf_clientes(fecha_inicial=None, fecha_final=None, formato='pdf'):
    if HTML is None:
        return HttpResponse("Error 501: La conversión a PDF binario (WeasyPrint) no está disponible.", status=501)

    clientes = _filtrar_clientes(fecha_inicial, fecha_final)

    context = {
        'titulo': 'Reporte de Clientes',
        'fecha_generacion': datetime.now().strftime('%d/%m/%Y %H:%M:%S'),
        'clientes': clientes,
        'filtro_fecha_inicial': fecha_inicial,
        'filtro_fecha_final': fecha_final,
    }

    try:
        html_string = render_to_string('reportes/reporte_clientes_pdf.html', context)
        html_doc = HTML(string=html_string)
        pdf_file = html_doc.write_pdf()

        response = HttpResponse(pdf_file, content_type='application/pdf')
        filename = f"reporte_clientes_{datetime.now().strftime('%Y%m%d_%H%M%S')}.pdf"
        
        disposition = 'inline' if formato in ['pdf_imprimir', 'inline'] else 'attachment'
        response['Content-Disposition'] = f'{disposition}; filename="{filename}"'
            
        return response
    except TemplateDoesNotExist:
        logger.error("Error fatal: No se encontró la plantilla 'reportes/reporte_clientes_pdf.html'")
        return HttpResponse("Error 500: No se encontró la plantilla 'reportes/reporte_clientes_pdf.html'.", status=500)
    except Exception as e:
        logger.error(f"Error al generar PDF de Clientes: {e}")
        return HttpResponse(f"Error al generar PDF de Clientes: {e}", status=500)


def generar_pdf_servicios(formato='pdf'):
    if HTML is None:
        return HttpResponse("Error 501: La conversión a PDF binario (WeasyPrint) no está disponible.", status=501)

    servicios = _filtrar_servicios()

    context = {
        'titulo': 'Reporte de Servicios',
        'fecha_generacion': datetime.now().strftime('%d/%m/%Y %H:%M:%S'),
        'servicios': servicios,
    }

    try: 
        html_string = render_to_string('reportes/reporte_servicios_pdf.html', context) 
        
        html_doc = HTML(string=html_string)
        pdf_file = html_doc.write_pdf()

        response = HttpResponse(pdf_file, content_type='application/pdf')
        filename = f"reporte_servicios_{datetime.now().strftime('%Y%m%d_%H%M%S')}.pdf"
        
        disposition = 'inline' if formato in ['pdf_imprimir', 'inline'] else 'attachment'

        response['Content-Disposition'] = f'{disposition}; filename="{filename}"'
        return response
    except TemplateDoesNotExist:
        logger.error("Error fatal: No se encontró la plantilla 'reportes/reporte_servicios_pdf.html'")
        return HttpResponse("Error 500: No se encontró la plantilla 'reportes/reporte_servicios_pdf.html'. Verifique la ruta y el nombre del archivo.", status=500)
    except AttributeError as e:
        logger.error(f"Error fatal generando reporte de servicios (Attr): {e}")
        return HttpResponse(f"Error interno al generar el reporte de Servicios. Detalle: {e}. Revise la plantilla HTML.", status=500)
    except Exception as e:
        logger.error(f"Error fatal generando reporte de servicios: {e}")
        return HttpResponse(f"Error al generar PDF de Servicios. Detalle: {e}", status=500)


def generar_pdf_especialistas(fecha_inicial=None, fecha_final=None, formato='pdf'):
    if HTML is None:
        return HttpResponse("Error 501: La conversión a PDF binario (WeasyPrint) no está disponible.", status=501)

    especialistas = _filtrar_especialistas(fecha_inicial, fecha_final)

    context = {
        'titulo': 'Reporte de Especialistas (Empleados)',
        'fecha_generacion': datetime.now().strftime('%d/%m/%Y %H:%M:%S'),
        'especialistas': especialistas,
        'filtro_fecha_inicial': fecha_inicial,
        'filtro_fecha_final': fecha_final,
    }

    try: 
        html_string = render_to_string('reportes/reporte_especialistas_pdf.html', context)
        
        html_doc = HTML(string=html_string)
        pdf_file = html_doc.write_pdf()

        response = HttpResponse(pdf_file, content_type='application/pdf')
        filename = f"reporte_especialistas_{datetime.now().strftime('%Y%m%d_%H%M%S')}.pdf"
        
        disposition = 'inline' if formato in ['pdf_imprimir', 'inline'] else 'attachment'

        response['Content-Disposition'] = f'{disposition}; filename="{filename}"'
        return response
    except TemplateDoesNotExist:
        logger.error("Error fatal: No se encontró la plantilla 'reportes/reporte_especialistas_pdf.html'")
        return HttpResponse("Error 500: No se encontró la plantilla 'reportes/reporte_especialistas_pdf.html'.", status=500)
    except AttributeError as e:
        logger.error(f"Error fatal generando reporte de especialistas (Attr): {e}")
        return HttpResponse(f"Error interno al generar el reporte de Especialistas. Detalle: {e}. Revise la plantilla 'reportes/reporte_especialistas_pdf.html'", status=500)
    except Exception as e:
        logger.error(f"Error fatal generando reporte de especialistas: {e}")
        return HttpResponse(f"Error al generar PDF de Especialistas. Detalle: {e}", status=500)

def generar_pdf_citas(fecha_inicial=None, fecha_final=None, formato='pdf'):
    if HTML is None:
        return HttpResponse("Error 501: La conversión a PDF binario (WeasyPrint) no está disponible.", status=501)

    citas_qs = _filtrar_citas(fecha_inicial, fecha_final)
    citas_procesadas = _procesar_citas(citas_qs)

    context = {
        'titulo': 'Reporte de Citas Agendadas',
        'fecha_generacion': datetime.now().strftime('%d/%m/%Y %H:%M:%S'),
        'citas': citas_procesadas,
        'filtro_fecha_inicial': fecha_inicial,
        'filtro_fecha_final': fecha_final,
    }

    try:
        html_string = render_to_string('reportes/reporte_citas_pdf.html', context)
        
        html_doc = HTML(string=html_string)
        pdf_file = html_doc.write_pdf()

        response = HttpResponse(pdf_file, content_type='application/pdf')
        filename = f"reporte_citas_{datetime.now().strftime('%Y%m%d_%H%M%S')}.pdf"
        
        disposition = 'inline' if formato in ['pdf_imprimir', 'inline'] else 'attachment'

        response['Content-Disposition'] = f'{disposition}; filename="{filename}"'
        return response
    except TemplateDoesNotExist:
        logger.error("Error fatal: No se encontró la plantilla 'reportes/reporte_citas_pdf.html'")
        return HttpResponse("Error 500: No se encontró la plantilla 'reportes/reporte_citas_pdf.html'.", status=500)
    except Exception as e:
        logger.error(f"Error fatal generando reporte de citas: {e}")
        return HttpResponse(f"Error al generar PDF de Citas. Detalle: {e}", status=500)